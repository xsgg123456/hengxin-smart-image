"""Pure extraction functions. Called only by the resource-limited worker."""
import csv
from io import BytesIO, StringIO
from zipfile import ZipFile

from defusedxml import ElementTree

MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_PAGES = 300
MAX_CELLS = 500_000


class DocumentError(ValueError):
    pass


class TextCollector:
    def __init__(self):
        self.parts = []
        self.size = 0

    def add(self, value):
        value = str(value).strip()
        if not value:
            return
        self.size += len(value.encode('utf-8')) + 1
        if self.size > MAX_TEXT_BYTES:
            raise DocumentError('解析文字超出 2 MiB 限制，请拆分文件')
        self.parts.append(value)

    def finish(self):
        if not self.parts:
            raise DocumentError('资料中没有可读取的文字或数值；扫描 PDF 请先转为文字版')
        return '\n'.join(self.parts)


def checked_archive(data, required):
    archive = ZipFile(BytesIO(data))
    try:
        entries = archive.infolist()
        if len(entries) > 2048 or sum(e.file_size for e in entries) > MAX_ARCHIVE_BYTES:
            raise DocumentError('文档解压大小或条目数量超出安全限制')
        names = [entry.filename for entry in entries]
        if len(names) != len(set(names)) or not required.issubset(names):
            raise DocumentError('文档结构损坏或与文件类型不一致')
        for entry in entries:
            if entry.flag_bits & 1:
                raise DocumentError('不支持加密文档，请解密后重新上传')
            if entry.file_size > max(entry.compress_size, 1) * 1000:
                raise DocumentError('文档压缩比例超出安全限制')
            if entry.filename.endswith(('.xml', '.rels')):
                # Reject entities/DTD in every XML part, including unused parts.
                ElementTree.fromstring(archive.read(entry), forbid_dtd=True)
        return archive
    except BaseException:
        archive.close()
        raise


def parse_docx(data, output):
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with checked_archive(data, {'[Content_Types].xml', 'word/document.xml'}) as archive:
        names = ['word/document.xml'] + sorted(
            name for name in archive.namelist()
            if name.startswith(('word/header', 'word/footer', 'word/footnotes',
                                'word/endnotes')) and name.endswith('.xml'))
        for name in names:
            root = ElementTree.fromstring(archive.read(name), forbid_dtd=True)
            for paragraph in root.iter(ns + 'p'):
                parts = []
                for node in paragraph.iter():
                    if node.tag == ns + 't':
                        parts.append(node.text or '')
                    elif node.tag == ns + 'tab':
                        parts.append('\t')
                    elif node.tag in {ns + 'br', ns + 'cr'}:
                        parts.append('\n')
                output.add(''.join(parts))


def parse_pdf(data, output):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise DocumentError('不支持加密 PDF，请解密后重新上传')
    if len(reader.pages) > MAX_PAGES:
        raise DocumentError('PDF 超过 300 页，请拆分文件')
    for page in reader.pages:
        output.add(page.extract_text() or '')


def parse_xlsx(data, output):
    from openpyxl import load_workbook

    with checked_archive(data, {'[Content_Types].xml', 'xl/workbook.xml'}):
        pass
    book = load_workbook(BytesIO(data), read_only=True, data_only=True,
                         keep_links=False, keep_vba=False)
    cells = 0
    try:
        for sheet in book.worksheets:
            # Do not trust producer dimensions: otherwise valid rows may be lost.
            sheet.reset_dimensions()
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_CELLS:
                    raise DocumentError('表格超过 500000 个单元格，请拆分文件')
                if any(value is not None and str(value).strip() for value in row):
                    output.add(f'[{sheet.title}] ' + '\t'.join(
                        '' if value is None else str(value) for value in row))
    finally:
        book.close()


def parse_xls(data, output):
    import xlrd

    book = xlrd.open_workbook(file_contents=data, on_demand=True,
                              logfile=StringIO(), ragged_rows=True)
    cells = 0
    try:
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            for row in sheet.get_rows():
                cells += len(row)
                if cells > MAX_CELLS:
                    raise DocumentError('表格超过 500000 个单元格，请拆分文件')
                values = [cell.value for cell in row]
                if any(str(value).strip() for value in values):
                    output.add(f'[{sheet.name}] ' + '\t'.join(map(str, values)))
            book.unload_sheet(index)
    finally:
        book.release_resources()


def parse_csv(data, output):
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            text = data.decode('gb18030')
        except UnicodeDecodeError:
            raise DocumentError('CSV 编码无效，请另存为 UTF-8 或 GB18030') from None
    if '\x00' in text:
        raise DocumentError('CSV 包含无效的二进制内容')
    csv.field_size_limit(MAX_TEXT_BYTES)
    cells = 0
    for row in csv.reader(StringIO(text, newline=''), strict=True):
        cells += len(row)
        if cells > MAX_CELLS:
            raise DocumentError('表格超过 500000 个单元格，请拆分文件')
        output.add('\t'.join(row))


def extract(data, extension):
    output = TextCollector()
    parsers = {'.pdf': parse_pdf, '.docx': parse_docx, '.xlsx': parse_xlsx,
               '.xls': parse_xls, '.csv': parse_csv}
    parsers[extension](data, output)
    return output.finish()
